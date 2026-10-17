from pathlib import Path
from typing import Union, List
import uuid

import docx

from .base_parser import BaseParser
from .models import ParsedDocument, DocumentSection, ParsedTable
from ..utils.logger import get_logger

logger = get_logger("docx_parser")


class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) documents."""

    def parse(self, file_path: Union[str, Path]) -> ParsedDocument:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {path}")

        file_hash = self.compute_sha256(path)
        file_size = path.stat().st_size
        doc_id = f"doc_{uuid.uuid4().hex[:10]}"

        doc = docx.Document(str(path))
        sections: List[DocumentSection] = []
        tables: List[ParsedTable] = []
        full_text_chunks: List[str] = []

        # Parse tables
        for tbl_idx, tbl in enumerate(doc.tables, start=1):
            rows_data = []
            headers = []
            for r_idx, row in enumerate(tbl.rows):
                row_vals = [cell.text.strip() for cell in row.cells]
                if r_idx == 0:
                    headers = row_vals
                else:
                    rows_data.append(row_vals)
            
            if headers or rows_data:
                table_obj = ParsedTable(
                    table_id=f"{doc_id}_tbl_{tbl_idx}",
                    sheet_or_page=f"Table {tbl_idx}",
                    headers=headers,
                    rows=rows_data,
                    title_or_caption=f"Embedded Table {tbl_idx}"
                )
                tables.append(table_obj)
                full_text_chunks.append(table_obj.to_markdown())

        # Parse paragraphs and section hierarchy
        current_title = "Document Overview"
        current_paras: List[str] = []
        current_level = 1

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                if current_paras:
                    sections.append(DocumentSection(
                        section_title=current_title,
                        content="\n\n".join(current_paras),
                        level=current_level
                    ))
                    current_paras = []
                current_title = text
                current_level = 1
            elif "heading 2" in style_name or "heading 3" in style_name:
                if current_paras:
                    sections.append(DocumentSection(
                        section_title=current_title,
                        content="\n\n".join(current_paras),
                        level=current_level
                    ))
                    current_paras = []
                current_title = text
                current_level = 2
            else:
                current_paras.append(text)
                full_text_chunks.append(text)

        if current_paras:
            sections.append(DocumentSection(
                section_title=current_title,
                content="\n\n".join(current_paras),
                level=current_level
            ))

        full_text = "\n\n".join(full_text_chunks)
        logger.info(f"Parsed DOCX '{path.name}' ({len(sections)} sections, {len(tables)} tables)")

        return ParsedDocument(
            document_id=doc_id,
            file_name=path.name,
            file_type="docx",
            file_size_bytes=file_size,
            content_hash=file_hash,
            total_pages_or_sheets=len(sections) or 1,
            sections=sections,
            tables=tables,
            full_text=full_text,
            metadata={"source": "Word Document"}
        )

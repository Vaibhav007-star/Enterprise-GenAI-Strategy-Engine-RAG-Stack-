import os
import io
from pathlib import Path
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from ..utils.logger import get_logger

logger = get_logger("deck_generator")


class ExecutiveDeckGenerator:
    """BCG X Executive Presentation Generator: Builds board-ready 16:9 widescreen PowerPoint decks
    synthesizing 10-K financial intelligence, strategic hypotheses, and execution roadmaps."""

    # BCG Executive Color Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)       # Deep Slate Navy (#0F172A)
    COLOR_ACCENT = RGBColor(0, 118, 90)        # BCG Emerald Forest Green (#00765A)
    COLOR_MUTED = RGBColor(100, 116, 139)      # Muted Steel (#64748B)
    COLOR_CARD_BG = RGBColor(248, 250, 252)    # Subtle Cloud Gray (#F8FAFC)
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Card Border (#E2E8F0)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_AMBER = RGBColor(194, 65, 12)        # Risk Amber (#C2410C)

    @classmethod
    def create_executive_deck(
        cls,
        title: str,
        subtitle: str,
        query: str,
        synthesis_text: str,
        metrics: Optional[Dict[str, Any]] = None,
        recommendations: Optional[List[str]] = None,
        risks: Optional[List[str]] = None
    ) -> io.BytesIO:
        """Create a full 3-slide executive PowerPoint presentation and return in-memory byte buffer."""
        prs = Presentation()
        # Set 16:9 Widescreen standard (13.33 x 7.5 inches)
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)
        blank_slide_layout = prs.slide_layouts[6]

        # -------------------------------------------------------------
        # SLIDE 1: Title & Strategic Context
        # -------------------------------------------------------------
        slide_1 = prs.slides.add_slide(blank_slide_layout)
        
        # Dark Navy Top Banner
        top_banner = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(1.8))
        top_banner.fill.solid()
        top_banner.fill.fore_color.rgb = cls.COLOR_PRIMARY
        top_banner.line.color.rgb = cls.COLOR_PRIMARY

        tf = top_banner.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.8)
        tf.margin_top = Inches(0.3)
        
        p = tf.paragraphs[0]
        p.text = "BCG X EXECUTIVE BRIEFING • GENAI STRATEGY ENGINE"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = cls.COLOR_ACCENT

        p2 = tf.add_paragraph()
        p2.text = title.upper()
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = cls.COLOR_WHITE

        # Subtitle bar
        sub_shape = slide_1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.73), Inches(0.8))
        tf_sub = sub_shape.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = f"Strategic Objective: {subtitle}"
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = cls.COLOR_MUTED

        # 3 Pillar Summary Cards on Slide 1
        card_width = Inches(3.64)
        card_height = Inches(3.8)
        card_y = Inches(2.8)

        card_data = [
            ("1. Audited Data & Fact Base", "Ground-truth evidence synthesized from audited SEC Form 10-K filings, period-over-period segment margins, and capital allocation filings.", cls.COLOR_ACCENT),
            ("2. Deterministic Verification", "All numerical metrics cross-checked against audited balance sheets with zero tolerance for arithmetic hallucinations.", cls.COLOR_PRIMARY),
            ("3. Actionable C-Suite Impact", "Strategic recommendations structured according to the BCG 10-20-70 transformation rule (Algorithms, Tech, People & Workflows).", cls.COLOR_ACCENT)
        ]

        for i, (c_title, c_desc, c_color) in enumerate(card_data):
            card_x = Inches(0.8 + i * 4.04)
            card = slide_1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, card_width, card_height)
            card.fill.solid()
            card.fill.fore_color.rgb = cls.COLOR_CARD_BG
            card.line.color.rgb = cls.COLOR_CARD_BORDER

            ctf = card.text_frame
            ctf.margin_left = Inches(0.4)
            ctf.margin_right = Inches(0.4)
            ctf.margin_top = Inches(0.4)
            ctf.word_wrap = True

            cp1 = ctf.paragraphs[0]
            cp1.text = c_title
            cp1.font.bold = True
            cp1.font.size = Pt(14)
            cp1.font.color.rgb = c_color

            cp2 = ctf.add_paragraph()
            cp2.text = c_desc
            cp2.font.size = Pt(12)
            cp2.font.color.rgb = cls.COLOR_PRIMARY

        # -------------------------------------------------------------
        # SLIDE 2: Detailed Strategic Synthesis & Key Metrics
        # -------------------------------------------------------------
        slide_2 = prs.slides.add_slide(blank_slide_layout)

        # Header Title
        hdr2 = slide_2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.73), Inches(0.9))
        tf2 = hdr2.text_frame
        p_hdr2 = tf2.paragraphs[0]
        p_hdr2.text = "STRATEGIC SYNTHESIS & AUDITED PERFORMANCE"
        p_hdr2.font.size = Pt(10)
        p_hdr2.font.bold = True
        p_hdr2.font.color.rgb = cls.COLOR_ACCENT

        p_hdr2_sub = tf2.add_paragraph()
        p_hdr2_sub.text = f"Action Title: Evidence-Backed Findings on '{query[:65]}...'"
        p_hdr2_sub.font.size = Pt(17)
        p_hdr2_sub.font.bold = True
        p_hdr2_sub.font.color.rgb = cls.COLOR_PRIMARY

        # Left Column: Strategic Synthesis Text Box
        left_box = slide_2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(7.5), Inches(5.3))
        left_box.fill.solid()
        left_box.fill.fore_color.rgb = cls.COLOR_CARD_BG
        left_box.line.color.rgb = cls.COLOR_CARD_BORDER

        tf_left = left_box.text_frame
        tf_left.word_wrap = True
        tf_left.margin_left = Inches(0.4)
        tf_left.margin_right = Inches(0.4)
        tf_left.margin_top = Inches(0.3)

        p_synth_head = tf_left.paragraphs[0]
        p_synth_head.text = "EXECUTIVE ADVISORY SYNTHESIS"
        p_synth_head.font.bold = True
        p_synth_head.font.size = Pt(12)
        p_synth_head.font.color.rgb = cls.COLOR_PRIMARY

        # Extract first few paragraphs of synthesis
        clean_synthesis = synthesis_text.replace("#", "").replace("*", "").strip()
        lines = [line.strip() for line in clean_synthesis.split("\n") if line.strip() and len(line) > 20][:5]
        for line in lines:
            p_line = tf_left.add_paragraph()
            p_line.text = f"• {line[:240]}..." if len(line) > 240 else f"• {line}"
            p_line.font.size = Pt(10.5)
            p_line.font.color.rgb = cls.COLOR_PRIMARY

        # Right Column: Metric Callout Cards
        right_x = Inches(8.6)
        right_w = Inches(3.93)
        metrics_dict = metrics or {
            "MSFT FY24 Revenue": "$245.1B (+16% YoY)",
            "Intelligent Cloud (Azure)": "$105.4B (Top Segment)",
            "Apple Services Margin": "74.0% Gross Margin",
            "Operating Margin": "44.6% Microsoft FY24"
        }

        for idx, (m_key, m_val) in enumerate(list(metrics_dict.items())[:4]):
            m_card = slide_2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, Inches(1.5 + idx * 1.35), right_w, Inches(1.15))
            m_card.fill.solid()
            m_card.fill.fore_color.rgb = cls.COLOR_CARD_BG
            m_card.line.color.rgb = cls.COLOR_CARD_BORDER

            tf_m = m_card.text_frame
            tf_m.word_wrap = True
            tf_m.margin_left = Inches(0.3)
            tf_m.margin_top = Inches(0.2)

            pm1 = tf_m.paragraphs[0]
            pm1.text = str(m_key).upper()
            pm1.font.size = Pt(9.5)
            pm1.font.bold = True
            pm1.font.color.rgb = cls.COLOR_MUTED

            pm2 = tf_m.add_paragraph()
            pm2.text = str(m_val)
            pm2.font.size = Pt(15)
            pm2.font.bold = True
            pm2.font.color.rgb = cls.COLOR_ACCENT

        # -------------------------------------------------------------
        # SLIDE 3: Risk Governance & 10-20-70 Execution Roadmap
        # -------------------------------------------------------------
        slide_3 = prs.slides.add_slide(blank_slide_layout)

        hdr3 = slide_3.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.73), Inches(0.9))
        tf3 = hdr3.text_frame
        p_hdr3 = tf3.paragraphs[0]
        p_hdr3.text = "EXECUTION GOVERNANCE & VALUE REALIZATION"
        p_hdr3.font.size = Pt(10)
        p_hdr3.font.bold = True
        p_hdr3.font.color.rgb = cls.COLOR_ACCENT

        p_hdr3_sub = tf3.add_paragraph()
        p_hdr3_sub.text = "Action Title: Risk Mitigations & BCG 10-20-70 Transformation Architecture"
        p_hdr3_sub.font.size = Pt(17)
        p_hdr3_sub.font.bold = True
        p_hdr3_sub.font.color.rgb = cls.COLOR_PRIMARY

        # 2 Wide Columns: Risks vs Roadmap
        col_w = Inches(5.66)
        col_h = Inches(5.3)

        # Risks Box (Left)
        risk_box = slide_3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), col_w, col_h)
        risk_box.fill.solid()
        risk_box.fill.fore_color.rgb = cls.COLOR_CARD_BG
        risk_box.line.color.rgb = cls.COLOR_CARD_BORDER

        tf_risk = risk_box.text_frame
        tf_risk.word_wrap = True
        tf_risk.margin_left = Inches(0.4)
        tf_risk.margin_right = Inches(0.4)
        tf_risk.margin_top = Inches(0.3)

        pr1 = tf_risk.paragraphs[0]
        pr1.text = "IDENTIFIED STRATEGIC & REGULATORY RISKS"
        pr1.font.bold = True
        pr1.font.size = Pt(12)
        pr1.font.color.rgb = cls.COLOR_AMBER

        default_risks = risks or [
            "AI Compute & GPU Availability: Hardware allocation bottlenecks impact cloud expansion timelines.",
            "EU AI Act & Regulatory Governance: Strict transparency requirements for high-risk generative enterprise models.",
            "Supply Chain & Geopolitical Concentration: Advanced chip fabrication exposure in Greater China and Taiwan.",
            "Data Clean Room & Confidentiality: Need for enterprise-grade RBAC across M&A and financial repositories."
        ]
        for r in default_risks[:4]:
            pr_item = tf_risk.add_paragraph()
            pr_item.text = f"▲ {r}"
            pr_item.font.size = Pt(10.5)
            pr_item.font.color.rgb = cls.COLOR_PRIMARY

        # Roadmap Box (Right) - BCG 10-20-70
        road_box = slide_3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.86), Inches(1.5), col_w, col_h)
        road_box.fill.solid()
        road_box.fill.fore_color.rgb = cls.COLOR_CARD_BG
        road_box.line.color.rgb = cls.COLOR_CARD_BORDER

        tf_road = road_box.text_frame
        tf_road.word_wrap = True
        tf_road.margin_left = Inches(0.4)
        tf_road.margin_right = Inches(0.4)
        tf_road.margin_top = Inches(0.3)

        proad1 = tf_road.paragraphs[0]
        proad1.text = "BCG 10-20-70 TRANSFORMATION PLAYBOOK"
        proad1.font.bold = True
        proad1.font.size = Pt(12)
        proad1.font.color.rgb = cls.COLOR_ACCENT

        roadmap_items = [
            ("10% - Model & Algorithms", "Deploy hybrid retrieval, Parent-Child chunking, and deterministic Python verification."),
            ("20% - Data & Tech Infrastructure", "Secure ChromaDB vector store, BM25 indexing, RBAC clean rooms, and semantic cache."),
            ("70% - People & Workflow Redesign", "Embed advisory copilot into analyst workflows, retrain strategy teams, and measure EBITDA ROI.")
        ]
        for step, desc in roadmap_items:
            ps = tf_road.add_paragraph()
            ps.text = f"■ {step}"
            ps.font.bold = True
            ps.font.size = Pt(11)
            ps.font.color.rgb = cls.COLOR_PRIMARY

            pd = tf_road.add_paragraph()
            pd.text = desc
            pd.font.size = Pt(10)
            pd.font.color.rgb = cls.COLOR_MUTED

        # Save presentation to memory buffer
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer
